# -*- coding: utf-8 -*-
"""零售经营决策系统 · 产品介绍 PPT（对外介绍，20 页）

设计约定：
- 16:9（13.333 × 7.5 in）
- 品牌色 #1E4A8F，深浅两级；强调 #B4703A；成功 #1E7A53
- 真实页面截图为主视觉；架构、闭环、部署为按 16:9 重绘的矢量图
- 每页一个结论：标题即观点，正文是证据
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
A = os.path.join(ROOT, "docs", "whitepaper-assets-v2", "doc")
B = os.path.join(ROOT, "docs", "whitepaper-assets-v2")
OUT = os.path.join(ROOT, "docs", "零售经营决策系统_产品介绍_V2.0.pptx")

NAVY = RGBColor(0x1E, 0x4A, 0x8F)
NAVY_D = RGBColor(0x0F, 0x2A, 0x4A)
NAVY_DD = RGBColor(0x14, 0x33, 0x57)
BLUE2 = RGBColor(0x2F, 0x55, 0x97)
BLUE_L = RGBColor(0x7F, 0xA8, 0xD9)
PALE = RGBColor(0xBF, 0xD6, 0xF2)
INK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x6D, 0x7C, 0x8F)
GRAY_L = RGBColor(0x9A, 0xAA, 0xBD)
BG_SOFT = RGBColor(0xF5, 0xF9, 0xFE)
BG_WARM = RGBColor(0xFF, 0xF9, 0xF3)
BG_GREEN = RGBColor(0xF0, 0xF9, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xB4, 0x70, 0x3A)
ORANGE_D = RGBColor(0x8A, 0x4B, 0x14)
ORANGE_T = RGBColor(0x8A, 0x73, 0x58)
GREEN = RGBColor(0x1E, 0x7A, 0x53)
GREEN_T = RGBColor(0x5F, 0x7A, 0x6C)
LINE = RGBColor(0xD5, 0xE2, 0xF2)
LINE2 = RGBColor(0xC9, 0xDA, 0xF0)
FONT = "微软雅黑"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.2, space_after=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for t, size, bold, color in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT
    return tb


def pic_fit(s, path, x, y, w, h, border=True, pad=0.03):
    im = Image.open(path)
    ratio = im.size[0] / im.size[1]
    if ratio > w / h:
        pw, ph = w, w / ratio
    else:
        ph, pw = h, h * ratio
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if border:
        rect(s, px - pad, py - pad, pw + 2 * pad, ph + 2 * pad, fill=WHITE,
             line=LINE, lw=1.0, radius=0.02)
    s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


def header(s, kicker, title, sub=None, tsize=25):
    rect(s, 0.62, 0.50, 0.09, 0.40, fill=NAVY, radius=0.4)
    text(s, 0.86, 0.45, 11.6, 0.32, [[(kicker, 12.5, True, NAVY)]])
    text(s, 0.86, 0.73, 11.9, 0.52, [[(title, tsize, True, NAVY_D)]])
    y = 1.34
    if sub:
        text(s, 0.86, 1.28, 11.9, 0.32, [[(sub, 12.5, False, GRAY)]])
        y = 1.70
    rect(s, 0.62, y - 0.09, 12.1, 0.018, fill=RGBColor(0xE3, 0xEA, 0xF3), radius=0)
    return y + 0.06


def footer(s, n):
    text(s, 0.62, 7.00, 8.0, 0.28,
         [[("零售经营决策 Agent · 产品介绍 V2.0", 9, False, GRAY_L)]])
    text(s, 12.0, 7.00, 0.75, 0.28, [[(str(n), 9.5, True, GRAY_L)]], align=PP_ALIGN.RIGHT)


def bullets(s, x, y, w, items, size=13.5, gap=0.62, dot=NAVY):
    for i, (head, rest) in enumerate(items):
        yy = y + i * gap
        rect(s, x, yy + 0.085, 0.10, 0.10, fill=dot, shape=MSO_SHAPE.OVAL)
        runs = []
        if head:
            runs.append((head, size, True, NAVY_D))
        runs.append((rest, size, False, INK))
        text(s, x + 0.24, yy - 0.02, w - 0.24, gap, [runs], line_spacing=1.18)


def side_panel(s, x, y, w, h, title, items, tcolor=NAVY_D, fill=BG_SOFT, lcolor=LINE2):
    rect(s, x, y, w, h, fill=fill, line=lcolor, lw=1.0, radius=0.07)
    text(s, x + 0.28, y + 0.24, w - 0.56, 0.34, [[(title, 13.5, True, tcolor)]])
    n = len(items)
    avail = h - 0.78
    step = avail / n
    for i, (h_, d) in enumerate(items):
        yy = y + 0.72 + i * step
        rect(s, x + 0.28, yy + 0.07, 0.085, 0.085, fill=NAVY, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.48, yy - 0.02, w - 0.76, 0.26, [[(h_, 11.5, True, NAVY_D)]])
        text(s, x + 0.48, yy + 0.22, w - 0.76, step - 0.24,
             [[(d, 9.5, False, GRAY)]], line_spacing=1.15)


# ================================================================ 1 封面
s = slide()
rect(s, 0, 0, W, H, fill=NAVY_D, radius=0)
rect(s, 0, 0, W, 0.13, fill=NAVY, radius=0)
rect(s, 0, H - 2.35, W, 2.35, fill=NAVY_DD, radius=0)
text(s, 1.1, 1.72, 11.0, 0.9, [[("Retail Decision Agent", 46, True, WHITE)]])
text(s, 1.1, 2.70, 11.0, 0.60, [[("零售经营决策系统 · 产品介绍", 29, True, PALE)]])
rect(s, 1.14, 3.52, 2.5, 0.045, fill=BLUE_L, radius=0)
text(s, 1.1, 3.82, 11.2, 0.55, [[
    ("自动发现异常 → 根因分析 → 补货 / 促销建议 → 授权范围内执行", 16.5, False, RGBColor(0xDA, 0xE7, 0xF7))]])
text(s, 1.1, H - 1.55, 6.5, 0.35, [[("版本 V2.0 ｜ 2026 年 9 月 30 日", 12.5, False, RGBColor(0x8F, 0xA8, 0xC4))]])
text(s, 1.1, H - 1.12, 11.2, 0.35, [[
    ("在线预览  https://shenyuanyuan77.github.io/retail-decision-agent/", 11, False, RGBColor(0x7F, 0x9C, 0xBC))]])

# ================================================================ 2 痛点
s = slide()
y = header(s, "问题", "经营异常每天都在发生，但发现它的人往往不是最该知道的人",
           "门店靠人巡检、靠经验判断，从异常发生到被处理，中间隔着好几层传递")
cards = [
    ("发现晚", "销量骤降往往在周会报表上\n才被看到，此时已经损失了\n数天销量"),
    ("归因难", "缺货、竞品、天气、客诉都能\n造成下跌，凭经验猜容易\n猜错方向"),
    ("动作慢", "分析完还要跨部门沟通，\n等补货到店，\n窗口期已经过去"),
    ("不可控", "就算让系统自动执行，\n谁来保证它\n不会做出越界动作"),
]
for i, (t, d) in enumerate(cards):
    x = 0.62 + i * 3.10
    rect(s, x, y + 0.10, 2.86, 2.60, fill=WHITE, line=LINE, lw=1.25, radius=0.08)
    rect(s, x, y + 0.10, 2.86, 0.07, fill=NAVY, radius=0)
    text(s, x + 0.24, y + 0.42, 2.4, 0.4, [[(t, 17, True, NAVY_D)]])
    text(s, x + 0.24, y + 0.98, 2.42, 1.5, [[(d, 11.5, False, GRAY)]], line_spacing=1.35)
rect(s, 0.62, y + 2.92, 12.1, 1.02, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.10)
text(s, 0.95, y + 3.16, 11.6, 0.6, [[
    ("传统数据看板只回答了「发生了什么」。", 15, True, NAVY_D),
    ("  真正缺的是另外三件事：为什么发生、该怎么办、办了之后效果如何。", 15, False, GRAY)]])
footer(s, 2)

# ================================================================ 3 产品定位
s = slide()
y = header(s, "产品定位", "把「发现 → 归因 → 建议 → 执行」做成一条自动闭环，并让每一步都可审计",
           "接入门店、销售、库存供应链、会员、客服与外部数据六大业务系统")
bullets(s, 0.62, y + 0.14, 6.9, [
    ("自动发现：", "持续扫描销售数据，主动发现销量异常，无需人工发起"),
    ("根因分析：", "对每个异常并行验证九类假设，输出带证据链与置信度的根因 Top3"),
    ("建议生成：", "按根因匹配打法库，生成补货或促销建议，含约束检查与成本收益估算"),
    ("授权执行：", "低风险动作自动执行，高风险动作强制转人工审批，全程留痕"),
], size=13.5, gap=0.68)
rect(s, 7.75, y + 0.12, 4.98, 2.80, fill=NAVY_D, radius=0.08)
text(s, 8.05, y + 0.38, 4.4, 0.36, [[("与传统看板的本质区别", 15, True, WHITE)]])
text(s, 8.05, y + 0.92, 4.4, 1.7, [[
    ("看板回答「发生了什么」", 12.5, False, PALE)],
    [("本产品回答「为什么、怎么办、效果如何」", 12.5, True, WHITE)],
    [("并把每一步交给可控可审计的 Agent 执行", 12.5, False, PALE)]], line_spacing=1.55)
rect(s, 0.62, y + 3.16, 12.1, 0.92, fill=BG_WARM, line=RGBColor(0xF3, 0xD9, 0xC0), lw=1.0, radius=0.10)
text(s, 0.95, y + 3.36, 11.6, 0.55, [[
    ("设计立场：", 14, True, ORANGE),
    ("系统不以「自动化」为唯一目标。所有高风险动作的判断权始终留在人手中，Agent 的服务身份硬上限为 L2。", 14, False, ORANGE_T)]])
footer(s, 3)

# ================================================================ 4 七步闭环
s = slide()
y = header(s, "工作方式", "一条异常从被发现到被验证改善，七个环节环环相扣",
           "目标 → 感知 → 诊断 → 决策 → 执行 → 治理 → 评估")
pic_fit(s, os.path.join(B, "slide-flow.png"), 0.62, y + 0.06, 12.1, 3.15)
rect(s, 0.62, y + 3.50, 5.94, 1.42, fill=BG_WARM, line=RGBColor(0xF3, 0xD9, 0xC0), lw=1.0, radius=0.09)
text(s, 0.90, y + 3.68, 5.4, 0.30, [[("治理不是事后合规，而是每一步的前置条件", 12.5, True, ORANGE_D)]])
text(s, 0.90, y + 4.02, 5.5, 0.85, [[
    ("写入前 Policy Engine 定风险等级；写入时幂等键 + dry-run 预演；写入后审计留痕、失败可补偿。", 10.5, False, ORANGE_T)]], line_spacing=1.3)
rect(s, 6.78, y + 3.50, 5.94, 1.42, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.09)
text(s, 7.06, y + 3.68, 5.4, 0.30, [[("每一步都可被验证，而不是被相信", 12.5, True, NAVY_D)]])
text(s, 7.06, y + 4.02, 5.5, 0.85, [[
    ("感知召回率 100%（线 90%）· 根因 Top3 命中率 100%（线 80%）· 约束违反率 0 · 自动执行成功率 100% · 高风险审批合规率 100%", 10.5, False, GRAY)]], line_spacing=1.3)
footer(s, 4)

# ================================================================ 5 架构
s = slide()
y = header(s, "系统架构", "事实来自业务系统，判断来自 Agent，控制来自治理层",
           "单体服务、分域设计：表与路由的边界即微服务拆分线")
pic_fit(s, os.path.join(B, "slide-arch.png"), 0.62, y + 0.06, 12.1, 4.34)
footer(s, 5)

# ================================================================ 6 能力总览
s = slide()
y = header(s, "能力总览", "五项核心能力，全部在真实系统中跑通并有对应测试",
           "数据来源：本机演示系统实跑（50 店 / 500 SKU / 180 天，随机种子 42）")
caps = [
    ("R1", "多系统接入", "六大 Mock 业务系统，\n统一 REST / OpenAPI 接入", "接口实测全部返回真实数据"),
    ("R2", "自动发现异常", "统计检测主动扫描，\n三类注入异常召回 100%", "3/3 故事线检出，高于 90% 验收线"),
    ("R3", "分析原因", "九类假设并行验证，\n根因 Top3 命中率 100%", "每根因 ≥2 条证据，高于 80% 验收线"),
    ("R4", "补货 / 促销建议", "根因 × 打法库生成建议，\n约束违反率 0", "MOQ / 毛利 / 预算全通过"),
    ("R5", "授权范围内执行", "低风险自动、高风险审批，\n审批合规率 100%", "无一例绕过人工审批"),
]
for i, (code, t, d, ev) in enumerate(caps):
    x = 0.62 + i * 2.47
    rect(s, x, y + 0.20, 2.26, 3.16, fill=WHITE, line=LINE, lw=1.25, radius=0.08)
    rect(s, x + 0.22, y + 0.46, 0.60, 0.32, fill=NAVY, radius=0.3)
    text(s, x + 0.22, y + 0.495, 0.60, 0.28, [[(code, 12, True, WHITE)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.22, y + 0.96, 1.9, 0.44, [[(t, 14, True, NAVY_D)]])
    text(s, x + 0.22, y + 1.54, 1.92, 1.06, [[(d, 10.5, False, GRAY)]], line_spacing=1.3)
    rect(s, x + 0.22, y + 2.62, 1.84, 0.60, fill=BG_GREEN, radius=0.12)
    text(s, x + 0.34, y + 2.74, 1.62, 0.44, [[(ev, 9, True, GREEN)]], line_spacing=1.2)
for i, (v, l) in enumerate([("42", "测试项全部通过"), ("100%", "异常召回率"), ("100%", "根因命中率"), ("0", "约束违反率")]):
    x = 0.62 + i * 3.10
    rect(s, x, y + 3.60, 2.86, 0.94, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.10)
    text(s, x + 0.26, y + 3.76, 1.35, 0.5, [[(v, 21, True, NAVY)]])
    text(s, x + 1.68, y + 3.87, 1.15, 0.38, [[(l, 11, False, GRAY)]])
footer(s, 6)

# ================================================================ 7 自动发现
s = slide()
y = header(s, "能力 R2 · 自动发现", "不用人工发起：系统自己扫数据，把异常按严重度排队",
           "前端标注「系统自动扫描销售数据发现异常，无需人工发起」")
pic_fit(s, os.path.join(A, "shot-anomalies-view.png"), 0.62, y + 0.02, 7.9, 4.32)
side_panel(s, 8.72, y + 0.02, 4.00, 4.32, "这一页在说什么", [
    ("严重度分级", "按偏差幅度自动划档：critical / high / medium / low"),
    ("两类异常", "销量骤降与骤升为主，另有门店客诉率激增"),
    ("观察 / 期望", "实测日均销量与基线中位数的对照，偏差列直接给出百分比"),
    ("时间窗口", "异常段的起止日期，便于回溯当时发生了什么"),
    ("状态流转", "open → diagnosed → decided，一眼看清处理到哪一步"),
])
footer(s, 7)

# ================================================================ 8 检测算法
s = slide()
y = header(s, "能力 R2 · 检测算法", "用统计方法而不是大模型猜测来发现异常，可解释、可复现",
           "传统算法负责发现与计算，大模型只负责叙述 —— 各做擅长的事")
steps = [
    ("建立基线", "取前 28 天数据，用滚动\n中位数与 MAD（中位数\n绝对偏差）刻画正常波动", "基线窗口 28 天"),
    ("逐日打分", "对近 14 天每天计算\nz-score 与偏差比例，\n量化当天偏离程度", "观测窗口 14 天"),
    ("识别异常段", "连续 ≥2 天偏离即认定为\n异常段，允许 1 天噪声\n桥接，避免切碎", "连续阈值 ≥2 天"),
    ("输出证据", "基线参数、逐日观测值与\n库存快照一并落库，\n每条异常都可逐条核对", "证据随异常同存"),
]
for i, (t, d, tag) in enumerate(steps):
    x = 0.62 + i * 3.10
    rect(s, x, y + 0.18, 2.86, 2.72, fill=WHITE, line=RGBColor(0x9F, 0xBE, 0xE3), lw=1.5, radius=0.08)
    rect(s, x, y + 0.18, 2.86, 0.07, fill=NAVY, radius=0)
    text(s, x + 0.26, y + 0.48, 2.4, 0.38, [[(t, 15.5, True, NAVY_D)]])
    text(s, x + 0.26, y + 1.02, 2.42, 1.30, [[(d, 11.5, False, GRAY)]], line_spacing=1.38)
    rect(s, x + 0.26, y + 2.16, 2.34, 0.46, fill=BG_SOFT, radius=0.12)
    text(s, x + 0.26, y + 2.30, 2.34, 0.30, [[(tag, 10.5, True, NAVY)]], align=PP_ALIGN.CENTER)
    if i < 3:
        text(s, x + 2.88, y + 1.32, 0.24, 0.4, [[("➜", 15, True, BLUE_L)]], align=PP_ALIGN.CENTER)
rect(s, 0.62, y + 3.14, 12.1, 0.86, fill=BG_GREEN, line=RGBColor(0xCB, 0xE7, 0xD8), lw=1.0, radius=0.10)
text(s, 0.95, y + 3.30, 11.6, 0.56, [[
    ("为什么不上来就用大模型：", 13, True, GREEN),
    ("异常发现要求可复现、可解释、可追责 —— 同样的数据必须给出同样的结论，因此感知层是纯统计的确定性算法。", 13, False, GREEN_T)]])
footer(s, 8)

# ================================================================ 9 根因分析
s = slide()
y = header(s, "能力 R3 · 根因分析", "对每个异常并行验证九类假设，输出带证据的根因 Top3",
           "每个根因至少 2 条证据；证据不足时不强行给结论，而是转人工分析")
pic_fit(s, os.path.join(A, "fig-detail-evidence.png"), 0.62, y + 0.02, 7.45, 4.32)
bx = 8.28
rect(s, bx, y + 0.02, 4.44, 4.32, fill=WHITE, line=LINE, lw=1.25, radius=0.07)
text(s, bx + 0.28, y + 0.24, 3.9, 0.34, [[("九类假设", 13.5, True, NAVY_D)]])
hyps = ["缺货", "供应链延迟", "自有促销拉动", "竞品促销分流", "天气影响",
        "客诉 / 口碑恶化", "临期未处置", "会员流失", "价格变化"]
for i, h in enumerate(hyps):
    col, row = i % 2, i // 2
    x = bx + 0.28 + col * 2.02
    yy = y + 0.74 + row * 0.42
    rect(s, x, yy, 1.86, 0.32, fill=BG_SOFT, line=LINE2, lw=0.75, radius=0.25)
    text(s, x + 0.10, yy + 0.045, 1.70, 0.26, [[(h, 10, False, NAVY_D)]])
rect(s, bx + 0.28, y + 2.76, 3.88, 0.014, fill=RGBColor(0xE3, 0xEA, 0xF3), radius=0)
text(s, bx + 0.28, y + 2.94, 3.9, 1.2, [[
    ("证据怎么来：", 11.5, True, NAVY_D),
    ("每类假设只调用只读工具取数（库存、单据、工单、促销、天气、竞品、复购），不直接读底层表。工具返回什么，Agent 就看到什么。", 10.5, False, GRAY)]], line_spacing=1.3)
footer(s, 9)

# ================================================================ 10 建议
s = slide()
y = header(s, "能力 R4 · 补货与促销建议", "建议卡不是一句话，而是带参数、约束与 ROI 的可执行方案",
           "先建议后执行：所有动作都以建议卡形式产生，人能看到它要做什么、值不值")
pic_fit(s, os.path.join(A, "shot-decisions.png"), 0.62, y + 0.02, 7.75, 4.32)
side_panel(s, 8.57, y + 0.02, 4.15, 4.32, "每张建议卡包含", [
    ("动作与参数", "补货量、折扣、预算、起止日期等完整 JSON 参数"),
    ("风险等级", "low / high，以及是否需要人工审批"),
    ("预估成本与收益", "补货按订单额约 5% 估算物流与资金成本"),
    ("ROI 与销量恢复", "收益成本比，以及预计能恢复多少销量"),
    ("约束检查明细", "覆盖天数、MOQ 倍数、毛利底线、折扣与预算上限逐项打勾"),
])
footer(s, 10)

# ================================================================ 11 执行与治理
s = slide()
y = header(s, "能力 R5 · 授权范围内执行", "低风险自动执行，高风险强制转人工，拒绝的动作同样留痕",
           "边界不是靠自觉，而是写死在策略里；Agent 服务身份硬上限 L2")
bullets(s, 0.62, y + 0.16, 6.5, [
    ("低风险 → 自动执行：", "常规补货、10% 以内折扣、店内客服回访任务，无需人工介入"),
    ("高风险 → 转人工审批：", "折扣 >10%、预算 >1 万、任何价格变更、跨区域调拨"),
    ("超上限 → 直接拒绝：", "折扣 >30%、预算 >10 万、白名单外的动作，策略层拦截"),
    ("执行时二次校验：", "执行前 Policy Engine 再校验一次，应对库存等状态变化"),
], size=12.5, gap=0.62)
rect(s, 7.42, y + 0.12, 5.30, 2.72, fill=WHITE, line=RGBColor(0xEF, 0xCD, 0xA9), lw=1.5, radius=0.07)
text(s, 7.70, y + 0.34, 4.8, 0.32, [[("真实审批单长什么样", 13.5, True, ORANGE)]])
text(s, 7.70, y + 0.80, 4.8, 1.6, [[
    ("[promo] [根因:临期未处置] 临期折价出清，降低损耗（>10% 折扣需审批） risk=high", 10, False, GRAY)],
    [("风险原因：折扣超过 10% 必须审批；Agent 服务身份上限 L2，高风险动作必须转人工审批", 10, False, GRAY)]], line_spacing=1.35)
rect(s, 7.70, y + 2.16, 2.45, 0.36, fill=BG_GREEN, radius=0.3)
text(s, 7.70, y + 2.225, 2.45, 0.28, [[("须由风控审批人（L3）处理", 10, True, GREEN)]], align=PP_ALIGN.CENTER)
for i, (v, l, sub) in enumerate([
    ("100%", "高风险审批合规率", "超标动作全部转人工，无一例绕过"),
    ("100%", "自动执行成功率", "低风险动作自动执行，全部成功"),
    ("0", "建议约束违反率", "约束不通过的候选直接丢弃"),
]):
    x = 0.62 + i * 4.10
    rect(s, x, y + 3.06, 3.90, 1.28, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.10)
    text(s, x + 0.24, y + 3.20, 3.4, 0.48, [[(v, 24, True, NAVY)]])
    text(s, x + 0.24, y + 3.72, 3.4, 0.28, [[(l, 11.5, True, NAVY_D)]])
    text(s, x + 0.24, y + 4.00, 3.5, 0.28, [[(sub, 9.5, False, GRAY)]])
footer(s, 11)

# ================================================================ 12 治理规则
s = slide()
y = header(s, "治理与安全", "自动执行、转人工、直接拒绝：三条通道各有明确边界",
           "Policy Engine 消费机器可读的规则库，调整规则不需要改代码")
pic_fit(s, os.path.join(B, "slide-gov.png"), 0.62, y + 0.06, 12.1, 3.34)
for i, (v, l, sub) in enumerate([
    ("L2", "Agent 服务身份硬上限", "可自主执行低风险动作，高风险一律转人工"),
    ("L3", "风控审批人唯一可批", "其他身份操作审批单会被拒绝并记入审计日志"),
    ("L4", "禁止动作策略层拦截", "不进入审批流程，避免把禁止项变成可协商项"),
]):
    x = 0.62 + i * 4.10
    rect(s, x, y + 3.60, 3.90, 1.42, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.10)
    text(s, x + 0.26, y + 3.76, 0.72, 0.42, [[(v, 21, True, NAVY)]])
    text(s, x + 1.06, y + 3.80, 2.6, 0.30, [[(l, 12, True, NAVY_D)]])
    text(s, x + 0.26, y + 4.30, 3.44, 0.56, [[(sub, 10.5, False, GRAY)]], line_spacing=1.25)
footer(s, 12)

# ================================================================ 13 审计
s = slide()
y = header(s, "治理与安全", "谁、在什么时候、对什么做了什么 —— 通过和拒绝都留痕",
           "按 trace_id 过滤，一键还原从发现到执行的完整链路")
pic_fit(s, os.path.join(A, "shot-audit-trace-view.png"), 0.62, y + 0.02, 7.9, 4.32)
side_panel(s, 8.72, y + 0.02, 4.00, 4.32, "审计日志记什么", [
    ("时间与主体", "区分 Agent 与人工操作者，标记每条动作的来源"),
    ("动作与资源", "扫描、诊断、决策、执行、审批等全部动作类型"),
    ("结果", "success / denied / error / dry_run 四种结果都记录"),
    ("详情", "完整请求与结果参数，可逐条展开核对"),
    ("全链路串联", "同一 trace_id 下，七层调用按时间倒序完整呈现"),
])
footer(s, 13)

# ================================================================ 14 四保险
s = slide()
y = header(s, "治理与安全", "写操作四保险：解决重复提交、误操作、无法追责、写错无法回退",
           "所有写接口统一经过一个受控写入入口，业务系统不暴露裸写通道")
four = [
    ("幂等", "相同幂等键 + 相同请求 → 重放首次结果，不重复下单；键相同但请求不同 → 直接拒绝", "重试与网络抖动导致重复下单"),
    ("预演", "任何写入可 dry-run，不落库但仍在审计日志留下一条 dry_run 记录", "误操作：先看预览再决定"),
    ("审计", "成功、拒绝、失败与预演均写审计日志，trace_id 贯穿全链路", "事后无法回答「谁做了什么」"),
    ("补偿", "补货、调拨、促销、工单均支持补偿回退，库存与单据同步恢复", "写错之后无法回退"),
]
for i, (t, d, s2) in enumerate(four):
    x = 0.62 + i * 3.10
    rect(s, x, y + 0.16, 2.86, 3.30, fill=WHITE, line=LINE, lw=1.25, radius=0.08)
    rect(s, x + 1.10, y + 0.44, 0.66, 0.66, fill=NAVY, shape=MSO_SHAPE.OVAL)
    text(s, x + 1.10, y + 0.60, 0.66, 0.4, [[(t, 16, True, WHITE)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.26, y + 1.32, 2.36, 1.3, [[(d, 11.5, False, GRAY)]], line_spacing=1.35)
    rect(s, x + 0.26, y + 2.62, 2.36, 0.62, fill=BG_WARM, radius=0.14)
    text(s, x + 0.40, y + 2.74, 2.10, 0.44, [[("解决：" + s2, 9.5, False, ORANGE)]], line_spacing=1.15)
footer(s, 14)

# ================================================================ 15 实证：执行记录
s = slide()
y = header(s, "实证", "跑一次 Agent：感知 → 诊断 → 决策 → 执行，四步结果全部落库",
           "本机演示系统实跑记录（业务日期 2026-10-07）")
pic_fit(s, os.path.join(A, "shot-executions.png"), 0.62, y + 0.02, 7.9, 4.32)
bx = 8.72
rect(s, bx, y + 0.02, 4.00, 4.32, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.07)
text(s, bx + 0.28, y + 0.26, 3.5, 0.34, [[("一次运行的结果", 13.5, True, NAVY_D)]])
res = [
    ("41", "条异常完成诊断", "九类假设逐条验证，证据不足者转人工"),
    ("3", "张建议卡生成", "含补货与促销两类动作，参数与 ROI 完整"),
    ("2", "条低风险自动执行", "在授权范围内直接落单，状态 success"),
    ("1", "条高风险转审批", "折扣超 10%，生成审批单等待人工决定"),
]
for i, (v, l, d) in enumerate(res):
    yy = y + 0.82 + i * 0.86
    text(s, bx + 0.28, yy - 0.04, 0.95, 0.42, [[(v, 19, True, NAVY)]])
    text(s, bx + 1.24, yy, 2.5, 0.26, [[(l, 11.5, True, NAVY_D)]])
    text(s, bx + 1.24, yy + 0.24, 2.54, 0.46, [[(d, 9.5, False, GRAY)]], line_spacing=1.15)
footer(s, 15)

# ================================================================ 16 审批实证
s = slide()
y = header(s, "实证", "待审批队列里躺着的，正是被规则拦下来的那一条",
           "审批历史与待审批队列同页呈现，状态、决定人与备注一目了然")
pic_fit(s, os.path.join(A, "shot-approvals.png"), 0.62, y + 0.02, 7.9, 4.32)
rect(s, 8.72, y + 0.02, 4.00, 4.32, fill=WHITE, line=RGBColor(0xEF, 0xCD, 0xA9), lw=1.25, radius=0.07)
text(s, 9.00, y + 0.26, 3.5, 0.34, [[("审批流的关键设计", 13.5, True, ORANGE)]])
pts = [
    ("Agent 无权自批", "服务身份上限 L2，L3 动作一律生成审批单，不存在自批路径"),
    ("只有 L3 能批", "风控审批人是唯一可处理审批单的角色，其他身份操作被拒绝并留痕"),
    ("决定回写决策卡", "审批结果同步更新决策卡状态，并记录决定人与备注"),
    ("通过后自动执行", "审批通过的决策由执行层携带审批单号执行，形成完整证据链"),
]
for i, (h_, d) in enumerate(pts):
    yy = y + 0.80 + i * 0.86
    text(s, 9.00, yy, 3.4, 0.26, [[(h_, 11.5, True, NAVY_D)]])
    text(s, 9.00, yy + 0.24, 3.44, 0.52, [[(d, 9.5, False, GRAY)]], line_spacing=1.15)
footer(s, 16)

# ================================================================ 17 技术实现
s = slide()
y = header(s, "技术实现", "四个关键设计决策，决定了这套系统好不好用、敢不敢用",
           "单体服务 + 分域设计，本地零配置可跑通全链路")
dec = [
    ("单库分域，而非一上来做微服务", "表与路由的边界就是微服务拆分线；本地零配置可跑，需要拆分时按域包一层即可"),
    ("规则 YAML 化，不写死在代码里", "指标、治理、根因、打法四类知识全部机器可读，业务方能自己维护"),
    ("大模型只做叙述，不参与计算", "全部数字来自工具返回的事实数据；叙述出现事实表外的数字，自动回退模板叙述"),
    ("trace_id 贯穿七层", "感知到评估全部带 trace，审计日志可一键还原案例的完整链路"),
]
for i, (h_, d) in enumerate(dec):
    x = 0.62 + (i % 2) * 6.30
    yy = y + 0.14 + (i // 2) * 1.94
    rect(s, x, yy, 5.82, 1.70, fill=WHITE, line=LINE, lw=1.25, radius=0.08)
    rect(s, x, yy, 0.07, 1.70, fill=NAVY, radius=0)
    text(s, x + 0.34, yy + 0.28, 5.2, 0.42, [[(h_, 14.5, True, NAVY_D)]])
    text(s, x + 0.34, yy + 0.84, 5.24, 0.74, [[(d, 11.5, False, GRAY)]], line_spacing=1.3)
rect(s, 0.62, y + 4.06, 12.1, 0.60, fill=BG_SOFT, line=LINE2, lw=1.0, radius=0.12)
text(s, 0.92, y + 4.21, 11.6, 0.36, [[
    ("前端原生 JS 无构建步骤 ｜ 后端 FastAPI + SQLModel + SQLite（WAL） ｜ Agent 为纯 Python 状态机 ｜ LLM 叙述层可选，未配置 API Key 时自动使用离线模板", 10.5, False, GRAY)]])
footer(s, 17)

# ================================================================ 18 部署与验收
s = slide()
y = header(s, "落地", "三种部署方式任选，验收指标全部达标",
           "本地零配置可跑通全链路；在线只读预览供外部查看")
pic_fit(s, os.path.join(B, "slide-deploy.png"), 0.62, y + 0.06, 12.1, 4.34)
footer(s, 18)

# ================================================================ 19 典型场景
s = slide()
y = header(s, "怎么用", "三条典型使用路径，从演示到日常巡检都能照做")
scen = [
    ("首次体验", "10 分钟跑通闭环", [
        "启动服务后打开控制台，进入经营总览",
        "点击【运行 Agent】，一次完成感知到执行",
        "在异常详情核对证据链与根因 Top3",
        "切到风控审批人，处理一条高风险审批",
    ], "看完就知道系统能做什么"),
    ("日常巡检", "每天 5 分钟过一遍", [
        "总览页看 KPI 卡与待处理计数",
        "异常中心按严重度处理 critical / high",
        "建议页复核 ROI，决定是否采纳",
        "审计页按 trace_id 回溯源链路",
    ], "把「看报表」变成「处理待办」"),
    ("对外演示", "用真实数据讲故事", [
        "先跑一次 Agent，让页面有真实数据",
        "用异常详情展示证据链的可核对性",
        "用审批中心演示治理边界",
        "用审计日志证明全过程可追溯",
    ], "不靠 PPT 讲，靠系统自己证明"),
]
for i, (tag, title, items, note) in enumerate(scen):
    x = 0.62 + i * 4.13
    rect(s, x, y + 0.14, 3.84, 4.34, fill=WHITE, line=LINE, lw=1.25, radius=0.08)
    rect(s, x, y + 0.14, 3.84, 0.07, fill=NAVY, radius=0)
    rect(s, x + 0.28, y + 0.40, 1.30, 0.32, fill=BG_SOFT, radius=0.3)
    text(s, x + 0.28, y + 0.445, 1.30, 0.28, [[(tag, 10.5, True, NAVY)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.28, y + 0.86, 3.3, 0.38, [[(title, 15.5, True, NAVY_D)]])
    for j, it in enumerate(items):
        yy = y + 1.42 + j * 0.52
        rect(s, x + 0.30, yy + 0.075, 0.09, 0.09, fill=BLUE_L, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.52, yy - 0.01, 3.06, 0.44, [[(it, 10.5, False, GRAY)]], line_spacing=1.15)
    rect(s, x + 0.28, y + 3.60, 3.28, 0.62, fill=BG_GREEN, radius=0.12)
    text(s, x + 0.42, y + 3.74, 3.0, 0.38, [[(note, 10, True, GREEN)]])
footer(s, 19)

# ================================================================ 20 结尾
s = slide()
rect(s, 0, 0, W, H, fill=NAVY_D, radius=0)
rect(s, 0, 0, W, 0.13, fill=NAVY, radius=0)
text(s, 1.2, 2.30, 11.0, 0.8, [[("让每一次经营异常，都有人接得住", 34, True, WHITE)]])
rect(s, 1.24, 3.36, 2.5, 0.045, fill=BLUE_L, radius=0)
text(s, 1.2, 3.66, 11.2, 0.7, [[
    ("自动发现 · 根因可查 · 建议可算 · 执行可控 · 全程可审计", 16, False, PALE)]])
for i, (k, v) in enumerate([
    ("在线预览", "https://shenyuanyuan77.github.io/retail-decision-agent/"),
    ("代码仓库", "https://github.com/shenyuanyuan77/retail-decision-agent"),
    ("本地启动", "scripts\\start.bat（Windows） 或 bash scripts/start.sh"),
]):
    yy = 5.24 + i * 0.52
    text(s, 1.2, yy, 1.5, 0.32, [[(k, 12.5, True, RGBColor(0x8F, 0xA8, 0xC4))]])
    text(s, 2.7, yy, 9.8, 0.32, [[(v, 12.5, False, RGBColor(0xDA, 0xE7, 0xF7))]])

prs.save(OUT)
print("已生成:", OUT)
print("页数:", len(prs.slides._sldIdLst))
